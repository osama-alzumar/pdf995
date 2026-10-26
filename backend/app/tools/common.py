"""Shared helpers: running external programs with a timeout, opening PDFs, checking outputs."""
from __future__ import annotations

import subprocess
from pathlib import Path

import pikepdf

from .. import errors
from ..config import settings


def run(cmd: list[str], timeout: int | None = None, ok_codes: tuple[int, ...] = (0,), env: dict | None = None) -> subprocess.CompletedProcess:
    try:
        proc = subprocess.run(cmd, capture_output=True, timeout=timeout or settings.job_timeout_seconds, env=env)
    except subprocess.TimeoutExpired:
        raise errors.timeout()
    except FileNotFoundError:
        raise errors.ToolError("engine_missing", "محرك المعالجة غير مثبت على الخادم.", "A processing engine is missing on the server.", 503)
    if proc.returncode not in ok_codes:
        raise errors.conversion_failed()
    return proc


def open_pdf(path: Path, password: str | None = None) -> pikepdf.Pdf:
    try:
        return pikepdf.open(path, password=password or "")
    except pikepdf.PasswordError:
        raise errors.wrong_password() if password else errors.password_required()
    except Exception:
        raise errors.cannot_open()


def is_pdf(path: Path) -> bool:
    with open(path, "rb") as fh:
        return b"%PDF-" in fh.read(1024)


def decrypted_copy(path: Path, password: str | None, work: Path) -> Path:
    """Tools that hand the file to other programs get an unencrypted copy."""
    with open_pdf(path, password) as pdf:
        if not pdf.is_encrypted:
            return path
        out = work / "decrypted.pdf"
        pdf.save(out)
        return out


# --- Output validation: a result is only returned after it has been re-opened and checked. ---

def validate_pdf(path: Path, expected_pages: int | None = None) -> int:
    if not path.exists() or path.stat().st_size == 0:
        raise errors.validation_failed()
    try:
        with pikepdf.open(path) as pdf:
            pages = len(pdf.pages)
    except Exception:
        raise errors.validation_failed()
    if pages == 0 or (expected_pages is not None and pages != expected_pages):
        raise errors.validation_failed()
    return pages


def validate_docx(path: Path) -> int:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception:
        raise errors.validation_failed()
    return len(document.paragraphs) + len(document.tables)


def validate_pptx(path: Path, expected_slides: int | None = None) -> int:
    from pptx import Presentation

    try:
        slides = len(Presentation(str(path)).slides)
    except Exception:
        raise errors.validation_failed()
    if slides == 0 or (expected_slides is not None and slides != expected_slides):
        raise errors.validation_failed()
    return slides


def validate_xlsx(path: Path) -> int:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True)
        sheets = len(workbook.sheetnames)
        workbook.close()
    except Exception:
        raise errors.validation_failed()
    if sheets == 0:
        raise errors.validation_failed()
    return sheets


def is_arabic(text: str) -> bool:
    letters = [c for c in text if c.isalpha()]
    if not letters:
        return False
    arabic = sum(1 for c in letters if "؀" <= c <= "ۿ" or "ﭐ" <= c <= "﻿")
    return arabic / len(letters) > 0.3
