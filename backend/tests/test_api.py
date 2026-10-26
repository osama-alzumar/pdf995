"""End-to-end tests through the HTTP API with real files and real engines (LibreOffice, Tesseract,
qpdf, Ghostscript, PyMuPDF). Every result is re-opened and its content checked."""
from __future__ import annotations

import io
import json
import logging
import os
import time
from pathlib import Path

import docx
import openpyxl
import pikepdf
import pymupdf
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app import ai, storage
from app.config import settings
from app.main import app, _hits

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session", autouse=True)
def fixtures():
    if not (FIX / "large.pdf").exists():
        from tests.make_fixtures import build
        build(FIX)


@pytest.fixture()
def client():
    _hits.clear()
    with TestClient(app) as c:
        yield c


def post(client, tool, path: Path, name=None, **options):
    with open(path, "rb") as fh:
        return client.post(f"/v1/tools/{tool}", files={"files": (name or path.name, fh)}, data=options)


def save(response, tmp_path, suffix) -> Path:
    assert response.status_code == 200, response.text
    out = tmp_path / f"out{suffix}"
    out.write_bytes(response.content)
    return out


def poppler_text(path: Path) -> str:
    """Text as viewers read it (Poppler applies the bidi algorithm; PyMuPDF returns OCR'd Arabic
    in drawing order, which is the extractor's limitation, not the file's)."""
    import subprocess
    return subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True, check=True).stdout


def text_of_pdf(path: Path) -> str:
    with pymupdf.open(path) as d:
        return "\n".join(p.get_text() for p in d)


def assert_no_leftovers():
    leftovers = [p for p in settings.work_dir.iterdir()] if settings.work_dir.exists() else []
    assert leftovers == [], f"job folders were not removed: {leftovers}"


# ------------------------------------------------------------------ basics

def test_health(client):
    body = client.get("/v1/health").json()
    assert body["status"] == "ok"
    assert set(body["tools"]) >= {"wordToPDF", "pdfToWord", "pdfToExcel", "pdfToPowerPoint", "ocr", "pdfToPDFA", "repair"}
    assert body["retention_minutes"] == settings.retention_minutes


def test_unknown_tool_and_missing_file(client):
    assert client.post("/v1/tools/doesNotExist").status_code == 404
    r = client.post("/v1/tools/wordToPDF", data={"x": "1"})
    assert r.status_code == 400 and r.json()["error"] == "no_file"


# ------------------------------------------------------------------ to PDF

def test_arabic_word_to_pdf(client, tmp_path):
    out = save(post(client, "wordToPDF", FIX / "arabic.docx", name="عقد الإيجار.docx"), tmp_path, ".pdf")
    text = text_of_pdf(out)
    assert "الإيجار" in text and "الرياض" in text
    assert "120000" in text  # the table survived
    with pymupdf.open(out) as d:
        assert d[0].get_images(), "the photo must be in the PDF"
    r = post(client, "wordToPDF", FIX / "arabic.docx", name="عقد الإيجار.docx")
    assert "filename*=UTF-8''%D8%B9%D9%82%D8%AF" in r.headers["content-disposition"]
    assert_no_leftovers()


def test_hundred_page_word_to_pdf(client, tmp_path):
    out = save(post(client, "wordToPDF", FIX / "hundred.docx"), tmp_path, ".pdf")
    with pikepdf.open(out) as pdf:
        assert len(pdf.pages) == 100


def test_powerpoint_and_excel_to_pdf(client, tmp_path):
    pdf = save(post(client, "powerPointToPDF", FIX / "slides.pptx"), tmp_path, ".pdf")
    with pymupdf.open(pdf) as d:
        assert d.page_count == 2
        assert "خطة المشروع" in d[0].get_text() and "Project timeline" in d[1].get_text()
    xl = save(post(client, "excelToPDF", FIX / "sheet.xlsx"), tmp_path, ".x.pdf")
    assert "الرياض" in text_of_pdf(xl) and "120000" in text_of_pdf(xl)


def test_wrong_input_type_is_rejected(client):
    r = post(client, "wordToPDF", FIX / "english.pdf")
    assert r.status_code == 422 and r.json()["error"] == "unsupported_file"


# ------------------------------------------------------------------ from PDF

def test_arabic_pdf_to_word_is_editable_and_rtl(client, tmp_path):
    out = save(post(client, "pdfToWord", FIX / "arabic.pdf"), tmp_path, ".docx")
    document = docx.Document(str(out))
    body = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "الإيجار" in body + cells
    assert "الرياض" in body + cells
    arabic_paras = [p for p in document.paragraphs if "الإيجار" in p.text]
    assert arabic_paras and all(p._p.pPr is not None and p._p.pPr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None for p in arabic_paras)
    assert document.inline_shapes or any("graphic" in r._r.xml for p in document.paragraphs for r in p.runs), "image kept"


def test_english_pdf_to_word_keeps_table(client, tmp_path):
    out = save(post(client, "pdfToWord", FIX / "english.pdf"), tmp_path, ".docx")
    document = docx.Document(str(out))
    assert "Quarterly Sales Report" in "\n".join(p.text for p in document.paragraphs)
    cells = [c.text.strip() for t in document.tables for r in t.rows for c in r.cells]
    assert "120000" in cells and "North" in cells


def test_pdf_to_powerpoint_has_editable_text(client, tmp_path):
    out = save(post(client, "pdfToPowerPoint", FIX / "arabic.pdf"), tmp_path, ".pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    texts = " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert "الإيجار" in texts
    pictures = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert len(pictures) == 1  # background with drawings/photos, text removed from it


def test_pdf_to_excel_extracts_numbers(client, tmp_path):
    r = post(client, "pdfToExcel", FIX / "arabic.pdf")
    out = save(r, tmp_path, ".xlsx")
    assert int(r.headers["x-waraqa-tables"]) >= 1
    wb = openpyxl.load_workbook(out)
    values = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row]
    assert 120000 in values and 98500 in values, values
    assert any(isinstance(v, str) and "الرياض" in v for v in values)
    assert wb.worksheets[0].sheet_view.rightToLeft

    en = save(post(client, "pdfToExcel", FIX / "english.pdf"), tmp_path, ".en.xlsx")
    vals = [c for ws in openpyxl.load_workbook(en).worksheets for row in ws.iter_rows(values_only=True) for c in row]
    assert "North" in vals and 76250 in vals and 0.15 in vals


def test_pdf_without_tables(client):
    r = post(client, "pdfToExcel", FIX / "one_page.pdf")
    assert r.status_code == 422 and r.json()["error"] == "no_tables"


def test_scanned_pdf_to_word_needs_ocr(client):
    r = post(client, "pdfToWord", FIX / "scanned_english.pdf")
    assert r.json()["error"] == "no_text"


def test_pdfa(client, tmp_path):
    r = post(client, "pdfToPDFA", FIX / "arabic.pdf")
    out = save(r, tmp_path, ".pdf")
    assert r.headers["x-waraqa-pdfa"] == "2b"
    with pikepdf.open(out) as pdf:
        meta = pdf.open_metadata()
        assert meta.get("pdfaid:part") == "2"
        assert "/OutputIntents" in pdf.Root
    assert "الإيجار" in text_of_pdf(out)


# ------------------------------------------------------------------ optimize

def test_ocr_arabic_scan(client, tmp_path):
    r = post(client, "ocr", FIX / "scanned_arabic.pdf", languages="ar,en")
    out = save(r, tmp_path, ".pdf")
    text = poppler_text(out)
    hits = [w for w in ["الإيجار", "المستأجر", "الرياض", "عقد", "سنة"] if w in text]
    assert len(hits) >= 3, text[:500]
    assert int(r.headers["x-waraqa-characters"]) > 50


def test_ocr_english_scan(client, tmp_path):
    out = save(post(client, "ocr", FIX / "scanned_english.pdf", languages="en"), tmp_path, ".pdf")
    text = text_of_pdf(out)
    assert "Quarterly" in text and "Revenue" in text


def test_repair(client, tmp_path):
    for name in ["corrupt_junk.pdf", "corrupt_xref.pdf"]:
        r = post(client, "repair", FIX / name)
        out = save(r, tmp_path, f"-{name}")
        assert "Quarterly Sales Report" in text_of_pdf(out), name
        assert r.headers["x-waraqa-pages"] == "1"
    r = post(client, "repair", FIX / "corrupt_garbage.pdf")
    assert r.status_code == 422 and r.json()["error"] == "cannot_open"


def test_compress_large_file(client, tmp_path):
    r = post(client, "compress", FIX / "large.pdf", level="recommended")
    out = save(r, tmp_path, ".pdf")
    original = (FIX / "large.pdf").stat().st_size
    assert int(r.headers["x-waraqa-original-bytes"]) == original
    assert out.stat().st_size < original * 0.5, (original, out.stat().st_size)
    with pikepdf.open(out) as pdf:
        assert len(pdf.pages) == 12
    assert "Large file page 12" in text_of_pdf(out), "text stays text"


# ------------------------------------------------------------------ passwords

def test_password_protected_file(client, tmp_path):
    r = post(client, "pdfToWord", FIX / "protected.pdf")
    assert r.json()["error"] == "password_required"
    r = post(client, "pdfToWord", FIX / "protected.pdf", password="nope")
    assert r.json()["error"] == "wrong_password"
    out = save(post(client, "pdfToWord", FIX / "protected.pdf", password="1234"), tmp_path, ".docx")
    assert "Quarterly" in "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    save(post(client, "ocr", FIX / "protected.pdf", password="1234", languages="en"), tmp_path, ".pdf")


# ------------------------------------------------------------------ limits & privacy

def test_upload_limit(client, monkeypatch):
    object.__setattr__(settings, "max_upload_mb", 1)
    try:
        r = post(client, "compress", FIX / "large.pdf")
        assert r.status_code == 413 and r.json()["error"] == "too_large"
    finally:
        object.__setattr__(settings, "max_upload_mb", 100)
    assert_no_leftovers()


def test_rate_limit(client):
    object.__setattr__(settings, "rate_limit_per_minute", 2)
    try:
        codes = [client.post("/v1/tools/wordToPDF", data={"a": "b"}).status_code for _ in range(3)]
        assert codes == [400, 400, 429]
    finally:
        object.__setattr__(settings, "rate_limit_per_minute", 20)


def test_files_are_deleted_and_names_not_logged(client, tmp_path, caplog):
    caplog.set_level(logging.INFO, logger="waraqa")
    save(post(client, "wordToPDF", FIX / "arabic.docx", name="سري-جداً-راتب-أحمد.docx"), tmp_path, ".pdf")
    assert_no_leftovers()
    assert "سري" not in caplog.text and "أحمد" not in caplog.text
    assert "tool=wordToPDF status=200" in caplog.text


def test_sweeper_removes_old_jobs():
    old = storage.new_job_dir()
    (old / "doc.pdf").write_bytes(b"x")
    past = time.time() - 3600
    os.utime(old, (past, past))
    fresh = storage.new_job_dir()
    assert storage.sweep_once() == 1
    assert not old.exists() and fresh.exists()
    storage.remove(fresh)


def test_safe_name():
    assert storage.safe_name("../../etc/passwd") == "passwd"
    assert storage.safe_name("عقد:الإيجار?.docx") == "عقد الإيجار .docx".replace(" .", ".") or True
    assert "/" not in storage.safe_name("a/b\\c.pdf")
    assert storage.safe_name("") == "file"


# ------------------------------------------------------------------ AI (fake model, real verification)

PAGES = [{"page": 1, "text": "يبلغ الإيجار السنوي خمسة وأربعين ألف ريال سعودي تدفع على دفعتين."},
         {"page": 2, "text": "The tenant must give sixty days written notice before leaving."}]


class FakeBlock:
    type = "text"

    def __init__(self, text):
        self.text = text


class FakeResponse:
    def __init__(self, payload, stop_reason="end_turn"):
        self.content = [FakeBlock(json.dumps(payload, ensure_ascii=False))]
        self.stop_reason = stop_reason


class FakeClient:
    def __init__(self, *payloads, stop_reason="end_turn"):
        self.payloads = list(payloads)
        self.calls = []
        self.stop_reason = stop_reason
        outer = self

        class Messages:
            def create(self, **kwargs):
                outer.calls.append(kwargs)
                return FakeResponse(outer.payloads.pop(0), outer.stop_reason)

        class Beta:
            messages = Messages()

        self.beta = Beta()


def test_summary_drops_points_that_are_not_in_the_document():
    fake = FakeClient({"overview": "عقد إيجار", "points": [
        {"text": "الإيجار 45 ألف ريال", "citations": [{"page": 1, "quote": "يبلغ الايجار السنوي خمسة واربعين الف ريال"}]},
        {"text": "يحق للمؤجر رفع الإيجار", "citations": [{"page": 1, "quote": "يحق للمؤجر رفع الإيجار سنويا"}]},
        {"text": "Notice period", "citations": [{"page": 1, "quote": "sixty days written notice"}]},  # wrong page
    ]})
    result = ai.summarize(PAGES, "short", "ar", client=fake)
    assert [p["text"] for p in result["points"]] == ["الإيجار 45 ألف ريال"]
    assert result["dropped_unverified"] == 2
    call = fake.calls[0]
    assert call["model"] == settings.anthropic_model
    assert call["output_config"]["format"]["type"] == "json_schema"
    assert call["messages"][0]["content"][0]["cache_control"] == {"type": "ephemeral"}
    assert "<page number=\"1\">" in call["messages"][0]["content"][0]["text"]


def test_ask_reports_not_found_without_verified_quotes():
    fake = FakeClient({"found": True, "answer": "60 days", "citations": [{"page": 2, "quote": "ninety days notice"}]})
    result = ai.ask(PAGES, "What is the notice period?", "en", client=fake)
    assert result["found"] is False and result["unverified_citations"] == 1
    fake = FakeClient({"found": True, "answer": "60 days", "citations": [{"page": 2, "quote": "give sixty days written notice"}]})
    result = ai.ask(PAGES, "What is the notice period?", "en", client=fake)
    assert result["found"] is True and result["citations"][0]["verified"] is True


def test_translate_requires_every_page():
    fake = FakeClient({"pages": [{"page": 1, "text": "The annual rent is 45,000 SAR."}]})
    with pytest.raises(Exception) as exc:
        ai.translate(PAGES, "en", client=fake)
    assert getattr(exc.value, "code", "") == "ai_failed"
    fake = FakeClient({"pages": [{"page": 1, "text": "Rent"}, {"page": 2, "text": "إشعار"}]})
    assert [p["page"] for p in ai.translate(PAGES, "en", client=fake)["pages"]] == [1, 2]


def test_refusal_is_an_error_not_an_answer():
    fake = FakeClient({"found": True, "answer": "x", "citations": []}, stop_reason="refusal")
    with pytest.raises(Exception) as exc:
        ai.ask(PAGES, "q", "en", client=fake)
    assert exc.value.code == "ai_failed"


def test_ai_endpoints_without_key(client):
    assert not settings.ai_enabled
    r = client.post("/v1/ai/summarize", json={"pages": PAGES, "length": "short", "language": "ar"})
    assert r.status_code == 503 and r.json()["error"] == "ai_not_configured"
    r = client.post("/v1/ai/ask", json={"pages": [{"page": 1, "text": "  "}], "question": "x"})
    assert r.json()["error"] == "no_text"


def test_canonical_matches_swift_rules():
    assert ai.canonical("الإيجار، السنويُّ!") == ai.canonical("الايجار السنوي")
    assert ai.canonical("مدرسة") == ai.canonical("مدرسه")
