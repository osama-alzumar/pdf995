"""Builds the real test documents used by the tests (and by the iOS team for manual checks).

Run: python -m tests.make_fixtures [out_dir]
"""
from __future__ import annotations

import io
import os
import random
import subprocess
import sys
from pathlib import Path

import docx
import openpyxl
import pikepdf
import pymupdf
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt

ARABIC_PARAS = [
    "عقد إيجار سكني",
    "تم الاتفاق بين المؤجر والمستأجر على تأجير الشقة الواقعة في مدينة الرياض لمدة سنة واحدة.",
    "يبلغ الإيجار السنوي خمسة وأربعين ألف ريال سعودي تدفع على دفعتين متساويتين.",
    "يلتزم المستأجر بالمحافظة على العين المؤجرة وإعادتها بحالة جيدة عند انتهاء العقد.",
]

ENGLISH_PARAS = [
    "Quarterly Sales Report",
    "This report summarizes the sales performance of the third quarter across all regions.",
    "Revenue grew by twelve percent compared with the previous quarter, driven by the eastern region.",
]

TABLE_AR = [["المنطقة", "المبيعات", "النمو"], ["الرياض", "120000", "12%"], ["جدة", "98500", "8%"], ["الدمام", "76250", "15%"]]
TABLE_EN = [["Region", "Sales", "Growth"], ["North", "120000", "12%"], ["South", "98500", "8%"], ["East", "76250", "15%"]]


def _rtl(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.insert(0, OxmlElement("w:bidi"))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in paragraph.runs:
        run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


def _photo(width=900, height=600, seed=1) -> bytes:
    rnd = random.Random(seed)
    img = Image.new("RGB", (width, height))
    draw = ImageDraw.Draw(img)
    for _ in range(300):
        x, y = rnd.randrange(width), rnd.randrange(height)
        draw.ellipse([x, y, x + rnd.randrange(20, 120), y + rnd.randrange(20, 120)],
                     fill=(rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)))
    noise = Image.effect_noise((width, height), 40).convert("RGB")
    img = Image.blend(img, noise, 0.25)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=92)
    return buf.getvalue()


def make_docx(path: Path, arabic: bool, with_table=True, with_image=True, pages: int = 1) -> None:
    document = docx.Document()
    paras = ARABIC_PARAS if arabic else ENGLISH_PARAS
    for n in range(pages):
        heading = document.add_heading(paras[0] + (f" {n + 1}" if pages > 1 else ""), level=1)
        if arabic:
            _rtl(heading)
        for text in paras[1:]:
            p = document.add_paragraph(text)
            if arabic:
                _rtl(p)
        if with_table and n == 0:
            rows = TABLE_AR if arabic else TABLE_EN
            table = document.add_table(rows=len(rows), cols=3)
            table.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = value
                    if arabic:
                        _rtl(cell.paragraphs[0])
            if arabic:
                tbl_pr = table._tbl.tblPr
                tbl_pr.append(OxmlElement("w:bidiVisual"))
        if with_image and n == 0:
            document.add_picture(io.BytesIO(_photo()), width=Inches(4))
        if n < pages - 1:
            document.add_page_break()
    for p in document.paragraphs:
        for run in p.runs:
            run.font.size = Pt(14)
            run.font.name = "Noto Naskh Arabic" if arabic else "Liberation Sans"
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), "Noto Naskh Arabic")
    document.save(str(path))


def soffice_pdf(source: Path, out_dir: Path) -> Path:
    out_dir = out_dir.resolve()
    profile = out_dir / ".lo-profile"
    subprocess.run(["soffice", "--headless", "--norestore", f"-env:UserInstallation=file://{profile}",
                    "--convert-to", "pdf", "--outdir", str(out_dir), str(source)],
                   check=True, capture_output=True, timeout=180, env=dict(os.environ, HOME=str(out_dir)))
    return out_dir / (source.stem + ".pdf")


def make_scanned(source_pdf: Path, out: Path, dpi=150) -> None:
    """Rasterizes every page: an image-only PDF with no text layer, like a scanner output."""
    src = pymupdf.open(source_pdf)
    dst = pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY)
        new = dst.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=pix.tobytes("jpeg", jpg_quality=80))
    dst.save(out)


def make_large(out: Path, pages=12) -> None:
    doc = pymupdf.open()
    for i in range(pages):
        page = doc.new_page()
        img = Image.open(io.BytesIO(_photo(1200, 1700, seed=i)))
        buf = io.BytesIO()
        img.save(buf, "PNG")  # lossless: big on purpose
        page.insert_image(page.rect, stream=buf.getvalue())
        page.insert_text((40, 40), f"Large file page {i + 1}", fontsize=14)
    doc.save(out, deflate=True)


def make_pptx(path: Path) -> None:
    prs = Presentation()
    for title, body in [("خطة المشروع", "المرحلة الأولى: التحليل"), ("Project timeline", "Phase two: build")]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    prs.save(str(path))


def make_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المبيعات"
    ws.sheet_view.rightToLeft = True
    for row in TABLE_AR:
        ws.append([int(v) if v.isdigit() else v for v in row])
    wb.save(str(path))


def build(out: Path) -> dict[str, Path]:
    out.mkdir(parents=True, exist_ok=True)
    f: dict[str, Path] = {}
    make_docx(out / "arabic.docx", arabic=True)
    make_docx(out / "english.docx", arabic=False)
    make_docx(out / "hundred.docx", arabic=False, with_table=False, with_image=False, pages=100)
    make_docx(out / "one_page.docx", arabic=False, with_table=False, with_image=False)
    make_pptx(out / "slides.pptx")
    make_xlsx(out / "sheet.xlsx")
    for name in ["arabic", "english", "hundred", "one_page"]:
        f[f"{name}_docx"] = out / f"{name}.docx"
        f[name] = soffice_pdf(out / f"{name}.docx", out)
    f["pptx"], f["xlsx"] = out / "slides.pptx", out / "sheet.xlsx"
    make_scanned(f["arabic"], out / "scanned_arabic.pdf", dpi=200)
    make_scanned(f["english"], out / "scanned_english.pdf", dpi=200)
    f["scanned_arabic"], f["scanned_english"] = out / "scanned_arabic.pdf", out / "scanned_english.pdf"
    with pikepdf.open(f["english"]) as pdf:
        pdf.save(out / "protected.pdf", encryption=pikepdf.Encryption(user="1234", owner="owner-5678", R=6))
    f["protected"] = out / "protected.pdf"
    data = f["english"].read_bytes()
    (out / "corrupt_junk.pdf").write_bytes(b"JUNK-BEFORE-HEADER\r\n" * 5 + data)
    xref = data.rfind(b"xref")
    (out / "corrupt_xref.pdf").write_bytes(data[:xref] + b"xref\n0 999\nbroken table\n" + data[xref + 40:])
    (out / "corrupt_garbage.pdf").write_bytes(bytes(random.Random(3).randrange(256) for _ in range(20000)))
    f["corrupt_junk"], f["corrupt_xref"], f["corrupt_garbage"] = (out / "corrupt_junk.pdf", out / "corrupt_xref.pdf",
                                                                  out / "corrupt_garbage.pdf")
    make_large(out / "large.pdf")
    f["large"] = out / "large.pdf"
    return f


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "tests/fixtures")
    for key, value in build(target).items():
        print(f"{key:16} {value.stat().st_size:>10,}  {value}")
